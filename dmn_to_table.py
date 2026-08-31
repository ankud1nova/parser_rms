import re
import sys
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.cell.cell import TYPE_STRING
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROJECT_FOLDER = Path(__file__).resolve().parent
DMN_FOLDER = PROJECT_FOLDER / "dmn_for_parser"
OUTPUT_FOLDER = PROJECT_FOLDER

EXPRESSION_TAGS = {
    "boxedContext",
    "context",
    "decisionTable",
    "filter",
    "for",
    "functionDefinition",
    "instanceOf",
    "invocation",
    "list",
    "literalExpression",
    "quantified",
    "relation",
    "some",
    "every",
}

DRG_TAGS = {
    "businessKnowledgeModel",
    "decision",
    "decisionService",
    "inputData",
    "knowledgeSource",
}

REFERENCE_TAGS = {
    "authorityRequirement",
    "informationRequirement",
    "knowledgeRequirement",
}


@dataclass
class ParseResult:
    file: Path
    sheet: str
    status: str
    source: str = ""
    version: str = ""
    decisions: int = 0
    decision_tables: int = 0
    rules: int = 0
    inputs: int = 0
    bkms: int = 0
    item_definitions: int = 0
    error: str = ""


def local_name(tag: str) -> str:
    if "}" in tag:
        return tag.split("}", 1)[1]
    return tag


def children(element: ET.Element, name: str) -> list[ET.Element]:
    return [node for node in list(element) if local_name(node.tag) == name]


def child(element: ET.Element, name: str) -> Optional[ET.Element]:
    found = children(element, name)
    return found[0] if found else None


def descendants(element: ET.Element, name: str) -> list[ET.Element]:
    return [node for node in element.iter() if local_name(node.tag) == name]


def first_descendant(element: ET.Element, name: str) -> Optional[ET.Element]:
    found = descendants(element, name)
    return found[0] if found else None


def attr(element: Optional[ET.Element], name: str, default: str = "") -> str:
    if element is None:
        return default
    return element.get(name) or default


def text_content(element: Optional[ET.Element]) -> str:
    if element is None:
        return ""
    return " ".join("".join(element.itertext()).split())


def dmn_text(element: Optional[ET.Element]) -> str:
    if element is None:
        return ""
    text_node = child(element, "text")
    if text_node is not None:
        return text_content(text_node)
    return text_content(element)


def direct_expression(element: ET.Element) -> Optional[ET.Element]:
    for node in list(element):
        if local_name(node.tag) in EXPRESSION_TAGS:
            return node
    return None


def expression_type(element: Optional[ET.Element]) -> str:
    if element is None:
        return ""
    return local_name(element.tag)


def xml_parse_file(dmn_file: Path) -> tuple[str, ET.Element]:
    xml_text = dmn_file.read_text(encoding="utf-8-sig", errors="ignore")
    try:
        root = ET.fromstring(xml_text)
    except ET.ParseError:
        start = xml_text.find("<")
        end = xml_text.rfind(">")
        if start < 0 or end <= start:
            raise
        root = ET.fromstring(xml_text[start:end + 1])
    return xml_text, root


def detect_dmn_source(xml_text: str) -> str:
    lower = xml_text.lower()
    sources = []
    if "kiegroup.org/dmn" in lower or "kie.org/dmn/extensions" in lower:
        sources.append("KIE / Drools / Red Hat / IBM BAMOE")
    if "camunda.org/schema/1.0/dmn" in lower or "camunda modeler" in lower:
        sources.append("Camunda")
    if "trisotech" in lower or "triso.tech" in lower:
        sources.append("Trisotech")
    if "signavio" in lower:
        sources.append("Signavio")
    return ", ".join(sources) if sources else "Unknown / Standard DMN"


def get_dmn_version(root: ET.Element) -> str:
    if root.tag.startswith("{"):
        namespace = root.tag.split("}", 1)[0].strip("{")
        match = re.search(r"/DMN/([^/]+)/MODEL", namespace, re.IGNORECASE)
        if match:
            return match.group(1)
        match = re.search(r"dmn/(\d+\.\d+)", namespace, re.IGNORECASE)
        if match:
            return match.group(1)
    return attr(root, "exporterVersion") or attr(root, "expressionLanguage")


def element_name(element: ET.Element, fallback: str) -> str:
    return attr(element, "name") or attr(element, "label") or attr(element, "id") or fallback


def variable_info(element: ET.Element) -> tuple[str, str]:
    variable = child(element, "variable")
    if variable is None:
        return "", ""
    return attr(variable, "name"), attr(variable, "typeRef")


def requirement_target(requirement: ET.Element) -> str:
    for ref_name in ("requiredDecision", "requiredInput", "requiredKnowledge", "requiredAuthority"):
        ref = child(requirement, ref_name)
        href = attr(ref, "href")
        if href:
            return href.lstrip("#")
    return attr(requirement, "href")


def collect_imports(root: ET.Element) -> list[list[str]]:
    rows = []
    for item in children(root, "import"):
        rows.append([
            attr(item, "name"),
            attr(item, "locationURI"),
            attr(item, "namespace"),
            attr(item, "importType"),
        ])
    return rows


def clean_sheet_name(name: str) -> str:
    name = re.sub(r"[\[\]\:\*\?\/\\]", "_", name or "Sheet").strip()
    return name[:31] or "Sheet"


def unique_sheet_name(wb: Workbook, desired_name: str) -> str:
    base = clean_sheet_name(desired_name)
    if base not in wb.sheetnames:
        return base
    counter = 1
    while True:
        suffix = f"_{counter}"
        candidate = clean_sheet_name(base[:31 - len(suffix)] + suffix)
        if candidate not in wb.sheetnames:
            return candidate
        counter += 1


def append_row(ws, values: list):
    ws.append(["" if value is None else value for value in values])


def append_section(ws, title: str):
    if ws.max_row > 1:
        append_row(ws, [])
    append_row(ws, [f"=== {title} ==="])


def append_key_value(ws, key: str, value: str):
    append_row(ws, [key, value])


def autosize_columns(ws):
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = str(cell.value) if cell.value is not None else ""
            longest_line = max(value.split("\n"), key=len, default="")
            max_length = max(max_length, len(longest_line))
        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 80)


def style_worksheet(ws):
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    section_fill = PatternFill("solid", fgColor="F2F2F2")
    error_fill = PatternFill("solid", fgColor="FCE4D6")
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows():
        first_value = row[0].value
        is_header = first_value in (
            "#",
            "Business name",
            "Decision",
            "Element",
            "File",
            "Input",
            "Item definition",
            "Rule #",
            "Source",
        )
        is_section = isinstance(first_value, str) and first_value.startswith("=== ")
        is_error = any(isinstance(cell.value, str) and cell.value.startswith("ERROR:") for cell in row)
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            if is_header:
                cell.font = Font(bold=True)
                cell.fill = header_fill
            if is_section:
                cell.font = Font(bold=True)
                cell.fill = section_fill
            if is_error:
                cell.fill = error_fill

    autosize_columns(ws)
    ws.freeze_panes = "A2"
    if ws.max_column > 1 and ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def clean_excel_string(value: str) -> str:
    return re.sub(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]", "", value)


def force_text_for_excel(wb: Workbook):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = clean_excel_string(cell.value)
                    cell.data_type = TYPE_STRING
                    cell.number_format = "@"


def input_header(input_node: ET.Element) -> str:
    label = element_name(input_node, "Input")
    input_expr = child(input_node, "inputExpression")
    expr_text = dmn_text(input_expr)
    type_ref = attr(input_expr, "typeRef") or attr(input_node, "typeRef")
    allowed_values = dmn_text(child(input_node, "inputValues"))
    parts = [label]
    if expr_text:
        parts.append(f"expr: {expr_text}")
    if type_ref:
        parts.append(f"type: {type_ref}")
    if allowed_values:
        parts.append(f"values: {allowed_values}")
    return "\n".join(parts)


def output_header(output_node: ET.Element) -> str:
    label = element_name(output_node, "Output")
    type_ref = attr(output_node, "typeRef")
    allowed_values = dmn_text(child(output_node, "outputValues"))
    parts = [label]
    if type_ref:
        parts.append(f"type: {type_ref}")
    if allowed_values:
        parts.append(f"values: {allowed_values}")
    return "\n".join(parts)


def write_decision_table(ws, decision: ET.Element, decision_table: ET.Element, start_row: int) -> tuple[int, int]:
    decision_name = element_name(decision, "Decision")
    inputs = children(decision_table, "input")
    outputs = children(decision_table, "output")
    rules = children(decision_table, "rule")
    row = start_row

    ws.cell(row=row, column=1, value=f"=== Decision table: {decision_name} ===")
    row += 1
    for key, value in [
        ("Decision", decision_name),
        ("Decision ID", attr(decision, "id")),
        ("Hit Policy", attr(decision_table, "hitPolicy")),
        ("Aggregation", attr(decision_table, "aggregation")),
        ("Preferred Orientation", attr(decision_table, "preferredOrientation")),
        ("Rules count", len(rules)),
    ]:
        ws.cell(row=row, column=1, value=key)
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1

    headers = ["Rule #"] + [input_header(node) for node in inputs] + [output_header(node) for node in outputs] + ["Annotation", "Rule ID"]
    for col, value in enumerate(headers, start=1):
        ws.cell(row=row, column=col, value=value)
    row += 1

    for rule_index, rule in enumerate(rules, start=1):
        input_entries = children(rule, "inputEntry")
        output_entries = children(rule, "outputEntry")
        annotation_entries = children(rule, "annotationEntry")
        values = [rule_index]
        for idx in range(len(inputs)):
            values.append(dmn_text(input_entries[idx]) if idx < len(input_entries) else "")
        for idx in range(len(outputs)):
            values.append(dmn_text(output_entries[idx]) if idx < len(output_entries) else "")
        values.append("; ".join(dmn_text(item) for item in annotation_entries if dmn_text(item)))
        values.append(attr(rule, "id"))
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        row += 1

    if not rules:
        ws.cell(row=row, column=1, value="Правила decisionTable не найдены")
        row += 1
    return row + 2, len(rules)


def write_expression(ws, owner: ET.Element, start_row: int, owner_label: str) -> int:
    name = element_name(owner, owner_label)
    expr = direct_expression(owner)
    row = start_row
    ws.cell(row=row, column=1, value=f"=== Expression: {name} ===")
    row += 1
    ws.cell(row=row, column=1, value=owner_label)
    ws.cell(row=row, column=2, value=name)
    row += 1
    ws.cell(row=row, column=1, value="Element ID")
    ws.cell(row=row, column=2, value=attr(owner, "id"))
    row += 1
    ws.cell(row=row, column=1, value="Expression Type")
    ws.cell(row=row, column=2, value=expression_type(expr) or "Unknown")
    row += 2

    if expr is None:
        ws.cell(row=row, column=1, value="Выражение не найдено")
        return row + 2

    expr_kind = expression_type(expr)
    if expr_kind == "context":
        ws.cell(row=row, column=1, value="#")
        ws.cell(row=row, column=2, value="Context entry")
        ws.cell(row=row, column=3, value="Expression type")
        ws.cell(row=row, column=4, value="Expression text")
        row += 1
        for idx, entry in enumerate(children(expr, "contextEntry"), start=1):
            variable_name, _ = variable_info(entry)
            inner_expression = next((node for node in list(entry) if local_name(node.tag) != "variable"), None)
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=variable_name)
            ws.cell(row=row, column=3, value=expression_type(inner_expression))
            ws.cell(row=row, column=4, value=dmn_text(inner_expression))
            row += 1
    elif expr_kind == "relation":
        columns = children(expr, "column")
        rows = children(expr, "row")
        headers = ["#"] + [element_name(column, "Column") for column in columns]
        for col, value in enumerate(headers, start=1):
            ws.cell(row=row, column=col, value=value)
        row += 1
        for idx, item in enumerate(rows, start=1):
            ws.cell(row=row, column=1, value=idx)
            for col, cell_node in enumerate(children(item, "literalExpression"), start=2):
                ws.cell(row=row, column=col, value=dmn_text(cell_node))
            row += 1
    elif expr_kind == "list":
        ws.cell(row=row, column=1, value="#")
        ws.cell(row=row, column=2, value="Item type")
        ws.cell(row=row, column=3, value="Item text")
        row += 1
        for idx, item in enumerate(list(expr), start=1):
            ws.cell(row=row, column=1, value=idx)
            ws.cell(row=row, column=2, value=local_name(item.tag))
            ws.cell(row=row, column=3, value=dmn_text(item))
            row += 1
    else:
        ws.cell(row=row, column=1, value="#")
        ws.cell(row=row, column=2, value="Expression type")
        ws.cell(row=row, column=3, value="Expression text")
        row += 1
        ws.cell(row=row, column=1, value=1)
        ws.cell(row=row, column=2, value=expr_kind)
        ws.cell(row=row, column=3, value=dmn_text(expr))
        row += 1
    return row + 2


def write_model_overview(ws, root: ET.Element, xml_text: str, dmn_file: Path):
    append_section(ws, "Паспорт DMN-файла")
    append_key_value(ws, "File", dmn_file.name)
    append_key_value(ws, "Relative path", str(dmn_file.relative_to(DMN_FOLDER)))
    append_key_value(ws, "Source", detect_dmn_source(xml_text))
    append_key_value(ws, "DMN Version", get_dmn_version(root))
    append_key_value(ws, "Definitions Name", attr(root, "name"))
    append_key_value(ws, "Definitions ID", attr(root, "id"))
    append_key_value(ws, "Namespace", attr(root, "namespace"))
    append_key_value(ws, "Exporter", attr(root, "exporter"))
    append_key_value(ws, "Exporter Version", attr(root, "exporterVersion"))

    imports = collect_imports(root)
    if imports:
        append_section(ws, "Imports")
        append_row(ws, ["Name", "Location URI", "Namespace", "Import type"])
        for row in imports:
            append_row(ws, row)

    append_section(ws, "DRG elements")
    append_row(ws, ["Element", "Business name", "ID", "Variable", "Type", "Expression", "Description"])
    for node in root.iter():
        kind = local_name(node.tag)
        if kind not in DRG_TAGS:
            continue
        variable_name, variable_type = variable_info(node)
        expr = direct_expression(node)
        append_row(ws, [
            kind,
            element_name(node, kind),
            attr(node, "id"),
            variable_name,
            variable_type,
            expression_type(expr),
            dmn_text(child(node, "description")),
        ])

    item_definitions = collect_item_definitions(root)
    if item_definitions:
        append_section(ws, "Item definitions")
        append_row(ws, ["Item definition", "Parent", "ID", "Type", "Collection", "Allowed values"])
        for row in item_definitions:
            append_row(ws, row)

    inputs = collect_inputs(root)
    if inputs:
        append_section(ws, "Inputs")
        append_row(ws, ["Input", "ID", "Variable", "Type", "Allowed values", "Description"])
        for row in inputs:
            append_row(ws, row)

    dependencies = collect_dependencies(root)
    if dependencies:
        append_section(ws, "Dependencies")
        append_row(ws, ["Source", "Source type", "Requirement", "Target ID"])
        for row in dependencies:
            append_row(ws, row)


def collect_inputs(root: ET.Element) -> list[list[str]]:
    rows = []
    for input_data in descendants(root, "inputData"):
        variable_name, variable_type = variable_info(input_data)
        rows.append([
            element_name(input_data, "Input"),
            attr(input_data, "id"),
            variable_name,
            variable_type,
            dmn_text(first_descendant(input_data, "allowedValues")),
            dmn_text(child(input_data, "description")),
        ])
    return rows


def collect_item_definitions(root: ET.Element) -> list[list[str]]:
    rows = []

    def walk_item(item: ET.Element, parent: str = ""):
        name = element_name(item, "Item")
        rows.append([
            name,
            parent,
            attr(item, "id"),
            attr(item, "typeRef"),
            attr(item, "isCollection"),
            dmn_text(child(item, "allowedValues")),
        ])
        for component in children(item, "itemComponent"):
            walk_item(component, name)

    for item_definition in children(root, "itemDefinition"):
        walk_item(item_definition)
    return rows


def collect_dependencies(root: ET.Element) -> list[list[str]]:
    rows = []
    for owner in root.iter():
        owner_kind = local_name(owner.tag)
        if owner_kind not in DRG_TAGS:
            continue
        for requirement in list(owner):
            req_kind = local_name(requirement.tag)
            if req_kind not in REFERENCE_TAGS:
                continue
            rows.append([
                element_name(owner, owner_kind),
                owner_kind,
                req_kind,
                requirement_target(requirement),
            ])
    return rows


def write_file_sheet(wb: Workbook, dmn_file: Path, xml_text: str, root: ET.Element) -> ParseResult:
    sheet_name = unique_sheet_name(wb, dmn_file.stem)
    ws = wb.create_sheet(sheet_name)
    write_model_overview(ws, root, xml_text, dmn_file)

    decisions = descendants(root, "decision")
    decision_tables_total = 0
    rules_total = 0
    row = ws.max_row + 2

    if not decisions:
        ws.cell(row=row, column=1, value="В файле не найдены decision. Проверьте листы Inputs, Types и Dependencies.")
    else:
        for decision in decisions:
            decision_tables = descendants(decision, "decisionTable")
            if decision_tables:
                decision_tables_total += len(decision_tables)
                for decision_table in decision_tables:
                    row, rules_count = write_decision_table(ws, decision, decision_table, row)
                    rules_total += rules_count
            else:
                row = write_expression(ws, decision, row, "Decision")

    bkms = descendants(root, "businessKnowledgeModel")
    for bkm in bkms:
        row = write_expression(ws, bkm, row, "Business knowledge model")

    return ParseResult(
        file=dmn_file,
        sheet=sheet_name,
        status="OK",
        source=detect_dmn_source(xml_text),
        version=get_dmn_version(root),
        decisions=len(decisions),
        decision_tables=decision_tables_total,
        rules=rules_total,
        inputs=len(descendants(root, "inputData")),
        bkms=len(bkms),
        item_definitions=len(descendants(root, "itemDefinition")),
    )


def write_summary_sheet(ws, results: list[ParseResult]):
    append_row(ws, [
        "File",
        "Sheet",
        "Status",
        "Source",
        "DMN Version",
        "Decisions",
        "Decision tables",
        "Rules",
        "Inputs",
        "BKM",
        "Item definitions",
        "Error",
    ])
    for result in results:
        append_row(ws, [
            str(result.file.relative_to(DMN_FOLDER)),
            result.sheet,
            result.status,
            result.source,
            result.version,
            result.decisions,
            result.decision_tables,
            result.rules,
            result.inputs,
            result.bkms,
            result.item_definitions,
            result.error,
        ])


def export_all_dmn_files():
    if not DMN_FOLDER.exists():
        raise FileNotFoundError(f"Папка не найдена: {DMN_FOLDER}")

    dmn_files = sorted(DMN_FOLDER.rglob("*.dmn"))
    if not dmn_files:
        raise ValueError(f"В папке нет .dmn файлов: {DMN_FOLDER}")

    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    output_file = OUTPUT_FOLDER / f"result_{timestamp}.xlsx"

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    inputs_ws = wb.create_sheet("Inputs")
    types_ws = wb.create_sheet("Types")
    deps_ws = wb.create_sheet("Dependencies")
    append_row(inputs_ws, ["Input", "ID", "Variable", "Type", "Allowed values", "Description", "File"])
    append_row(types_ws, ["Item definition", "Parent", "ID", "Type", "Collection", "Allowed values", "File"])
    append_row(deps_ws, ["Source", "Source type", "Requirement", "Target ID", "File"])

    results: list[ParseResult] = []
    for dmn_file in dmn_files:
        try:
            xml_text, root = xml_parse_file(dmn_file)
            results.append(write_file_sheet(wb, dmn_file, xml_text, root))

            relative_file = str(dmn_file.relative_to(DMN_FOLDER))
            for row in collect_inputs(root):
                append_row(inputs_ws, row + [relative_file])
            for row in collect_item_definitions(root):
                append_row(types_ws, row + [relative_file])
            for row in collect_dependencies(root):
                append_row(deps_ws, row + [relative_file])
        except Exception as error:
            results.append(ParseResult(
                file=dmn_file,
                sheet="",
                status="ERROR",
                error=f"ERROR: {type(error).__name__}: {error}",
            ))

    summary = wb.create_sheet("Summary", 0)
    write_summary_sheet(summary, results)

    for ws in wb.worksheets:
        style_worksheet(ws)

    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    force_text_for_excel(wb)
    wb.save(output_file)

    ok_count = sum(1 for result in results if result.status == "OK")
    error_count = len(results) - ok_count
    print("Готово.")
    print(f"DMN-файлов найдено: {len(dmn_files)}")
    print(f"Успешно обработано: {ok_count}")
    print(f"С ошибками: {error_count}")
    print(f"Excel-файл: {output_file}")


if __name__ == "__main__":
    try:
        export_all_dmn_files()
    except Exception as error:
        print(f"Ошибка: {error}", file=sys.stderr)
        sys.exit(1)
